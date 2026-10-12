# core/excel_export.py
import pandas as pd
from pathlib import Path
from openpyxl.styles import Font, PatternFill
from pandas import ExcelWriter

HIDE_COLS = [
    "VALUE_SCORE_STATIC", "QUALITY_SCORE_STATIC",
    "earnings_yield_boost", "ebit_boost", "fcf_cagr_boost", "revenue_boost",
    "roic_boost", "op_margin_boost",
    "trailing_earnings_yield", "forward_earnings_yield", "ev_ebit_yield",
    "ev_ebitda_yield", "fcf_yield", "sales_yield", "roic", "gpa",
    "accruals_ratio", "operating_profit_margin"
]

def export_to_excel(df: pd.DataFrame, path: Path):
    with ExcelWriter(path, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name="metrics")
        ws = writer.sheets["metrics"]

        # Format header
        dark_green = Font(bold=True, color="006400")
        light_green_fill = PatternFill("solid", fgColor="C6EFCE")
        dark_blue = Font(bold=True, color="00008B")
        light_blue_fill = PatternFill("solid", fgColor="CCECFF")
        purple = Font(bold=True, color="800080")
        bold = Font(bold=True)
        light_purple_fill = PatternFill("solid", fgColor="E6CCFF")

        # -- COLUMN GROUPS --
        group1 = [  # dark green font
            "trailing_earnings_yield","forward_earnings_yield","ev_ebit_yield","ev_ebitda_yield","fcf_yield",
            "sales_yield","roic","gpa","accruals_ratio","operating_profit_margin"
        ]
        group2 = [  # light green bg
            "trailing_earnings_yield pct","forward_earnings_yield pct","ev_ebit_yield pct","ev_ebitda_yield pct",
            "fcf_yield pct","sales_yield pct","roic pct","gpa pct","accruals_ratio pct","operating_profit_margin pct"
        ]
        group3 = [  # dark blue font
            "earnings_yield_boost","ebit_boost","fcf_cagr_boost","revenue_boost","roic_boost","op_margin_boost"
        ]
        group4 = [  # light blue bg
            "earnings_yield_boost_pct","ebit_boost_pct","fcf_cagr_boost_pct","revenue_boost_pct","roic_boost_pct","op_margin_boost_pct"
        ]
        group5 = ["HQM_SCORE", "DOLLAR_VOLUME", "TURNOVER_RATIO"]

        # Autofilter top row
        ws.auto_filter.ref = ws.dimensions

        for idx, col_name in enumerate(df.columns, 1):
            cell = ws.cell(row=1, column=idx)
            cell.font = bold
            if col_name in group1: cell.font = dark_green
            if col_name in group2: cell.fill = light_green_fill
            if col_name in group3: cell.font = dark_blue
            if col_name in group4: cell.fill = light_blue_fill
            if col_name in group5: cell.font = purple

            if col_name in HIDE_COLS:
                ws.column_dimensions[cell.column_letter].hidden = True

            max_len = max(df[col_name].astype(str).map(len).max(), len(col_name)) + 2
            ws.column_dimensions[cell.column_letter].width = max_len

    print(f"✅ Saved full formatted Excel file → {path}")
