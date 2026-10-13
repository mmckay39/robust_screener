# core/excel_export.py
import pandas as pd
from pathlib import Path
from openpyxl.styles import Font, PatternFill
from pandas import ExcelWriter

from calc.rank import VALUE_FACTORS, QUALITY_FACTORS, BOOST_COLS

# Columns written in scientific notation. The cell stays numeric, so it still works in formulas.
SCI_COLS = ["DOLLAR_VOLUME", "TURNOVER_RATIO"]
SCI_FORMAT = "0.000E+00"

HIDE_COLS = [
    "VALUE_SCORE_STATIC", "QUALITY_SCORE_STATIC",
    *BOOST_COLS, *VALUE_FACTORS, *QUALITY_FACTORS,
]

def export_to_excel(df: pd.DataFrame, path: Path):
    with ExcelWriter(path, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name="metrics")
        ws = writer.sheets["metrics"]

        # Format header
        dark_green = Font(bold=True, color="006400")
        light_green_fill = PatternFill("solid", fgColor="C6EFCE")
        dark_blue = Font(bold=True, color="00008B")
        light_blue_fill = PatternFill("solid", fgColor="CCECFF")
        purple = Font(bold=True, color="800080")
        bold = Font(bold=True)

        # -- COLUMN GROUPS (derived from calc/rank.py so the export and scoring can't drift apart) --
        group1 = [*VALUE_FACTORS, *QUALITY_FACTORS]                  # dark green font
        group2 = [f"{c} pct" for c in group1]                        # light green bg
        group3 = list(BOOST_COLS)                                    # dark blue font
        group4 = [f"{c}_pct" for c in BOOST_COLS]                    # light blue bg
        group5 = ["HQM_SCORE", "DOLLAR_VOLUME", "TURNOVER_RATIO"]

        # Autofilter top row
        ws.auto_filter.ref = ws.dimensions

        for idx, col_name in enumerate(df.columns, 1):
            cell = ws.cell(row=1, column=idx)
            cell.font = bold
            if col_name in group1: cell.font = dark_green
            if col_name in group2: cell.fill = light_green_fill
            if col_name in group3: cell.font = dark_blue
            if col_name in group4: cell.fill = light_blue_fill
            if col_name in group5: cell.font = purple

            if col_name in HIDE_COLS:
                ws.column_dimensions[cell.column_letter].hidden = True

            if col_name in SCI_COLS:
                for row in range(2, ws.max_row + 1):
                    ws.cell(row=row, column=idx).number_format = SCI_FORMAT

            max_len = max(df[col_name].map(lambda v: len(str(v))).max(), len(col_name)) + 2
            ws.column_dimensions[cell.column_letter].width = max_len

    print(f"✅ Saved full formatted Excel file → {path}")
